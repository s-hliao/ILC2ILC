#!/usr/bin/env python3
"""tables_to_ppt.py (2026-10-23): every table in media/{pogo,car,quad}/tables.md as (1) tables.pptx, one native PowerPoint
table per slide (copy the table into your deck), (2) tables.html (open in a browser, select a table, copy, paste into
PowerPoint: it pastes as a native table), (3) tables_tsv/*.tsv (tab-separated; pastes into a PowerPoint/Excel table).
    ~/miniconda3/envs/ilcmjx/bin/python tables_to_ppt.py
"""
import html
import os
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
DOMAINS = [("pogo", "Pogo"), ("car", "Car"), ("quad", "Quadruped")]
BLUE, GRAY, INK = RGBColor(0x25, 0x6A, 0xBF), RGBColor(0xEC, 0xEB, 0xE8), RGBColor(0x22, 0x22, 0x22)


def parse(md):
    """[(title, rows, note)] from a tables.md: a title line (### or **bold**) before each pipe table"""
    out, title, rows, note = [], "", [], []
    for ln in md.splitlines() + [""]:
        s = ln.strip().replace("|dv|", "speed err. dv")
        if s.startswith("|"):
            cells = [c.strip() for c in s.strip("|").split("|")]
            if not all(re.fullmatch(r":?-+:?", c) for c in cells if c):
                rows.append(cells)
            continue
        if rows:
            out.append([title, rows, ""])
            rows = []
        if s.startswith("###") or (s.startswith("**") and s.endswith("**")):
            title = re.sub(r"^(Pogo|Car|Quad)\s*", "", s.lstrip("#").strip().strip("*")).lstrip("(").strip()
            if title.count(")") > title.count("("):
                title = title.replace(")", "", 1)
        elif s and out:
            out[-1][2] = (out[-1][2] + " " + s).strip()
    return out


def is_ours(r):
    return "DistILC" in r[0] or r[0].startswith("ours +") or r[0].startswith("stage,")


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    H = ["<html><head><meta charset='utf-8'><style>body{font-family:Calibri,Arial;margin:24px}"
         "table{border-collapse:collapse;margin:8px 0 28px}td,th{border:1px solid #bbb;padding:4px 10px;font-size:14px}"
         "th{background:#ecebe8;text-align:left}tr.ours td{color:#256abf;font-weight:bold}p{color:#555;font-size:12px;max-width:1100px}"
         "</style></head><body>"]
    os.makedirs(f"{HERE}/tables_tsv", exist_ok=True)
    for dom, name in DOMAINS:
        for t, (title, rows, note) in enumerate(parse(open(f"{HERE}/{dom}/tables.md").read())):
            ncol = max(len(r) for r in rows)
            rows = [r + [""] * (ncol - len(r)) for r in rows]
            # tsv
            with open(f"{HERE}/tables_tsv/{dom}_{t + 1}.tsv", "w") as fh:
                fh.write("\n".join("\t".join(r) for r in rows) + "\n")
            # html
            H.append(f"<h3>{html.escape(name)}: {html.escape(title)}</h3><table>")
            for i, r in enumerate(rows):
                tag = "th" if i == 0 else "td"
                cls = " class='ours'" if i and is_ours(r) else ""
                H.append(f"<tr{cls}>" + "".join(f"<{tag}>{html.escape(c)}</{tag}>" for c in r) + "</tr>")
            H.append("</table>" + (f"<p>{html.escape(note)}</p>" if note else ""))
            # pptx: title + native table + footnote
            s = prs.slides.add_slide(prs.slide_layouts[6])
            tb = s.shapes.add_textbox(Inches(0.4), Inches(0.25), Inches(12.5), Inches(0.8)).text_frame
            tb.word_wrap = True
            tb.text = f"{name}: {title}"
            tb.paragraphs[0].runs[0].font.size = Pt(20)
            tb.paragraphs[0].runs[0].font.bold = True
            h = Inches(0.42) * len(rows)
            tab = s.shapes.add_table(len(rows), ncol, Inches(0.4), Inches(1.2), Inches(12.5), h).table
            for i, r in enumerate(rows):
                for j, c in enumerate(r):
                    cell = tab.cell(i, j)
                    cell.text = c
                    run = cell.text_frame.paragraphs[0].runs
                    if run:
                        run[0].font.size = Pt(13 if ncol <= 5 else 11)
                        run[0].font.bold = i == 0 or is_ours(r)
                        run[0].font.color.rgb = INK if i == 0 else (BLUE if is_ours(r) else INK)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = GRAY if i == 0 else RGBColor(0xFF, 0xFF, 0xFF)
            if note:
                nb = s.shapes.add_textbox(Inches(0.4), Inches(1.35) + h, Inches(12.5), Inches(1.2)).text_frame
                nb.word_wrap = True
                nb.text = note
                nb.paragraphs[0].runs[0].font.size = Pt(10)
    H.append("</body></html>")
    open(f"{HERE}/tables.html", "w").write("\n".join(H))
    prs.save(f"{HERE}/tables.pptx")
    print(f"{HERE}/tables.pptx", f"{HERE}/tables.html", f"{HERE}/tables_tsv/")


if __name__ == "__main__":
    main()
